import os
import sys
import json
import fitz  # PyMuPDF
from docx import Document

if sys.platform == "win32":
    sys.stdout.reconfigure(encoding="utf-8")
    sys.stderr.reconfigure(encoding="utf-8")

data_dir = "public/data"

print("======================================================================")
print("  PHASE 0: COMPREHENSIVE DATA INVENTORY & AUDIT")
print("======================================================================")

# 1. Inspect PDFs
pdf_files = [f for f in os.listdir(data_dir) if f.endswith(".pdf")]
for pdf_file in pdf_files:
    pdf_path = os.path.join(data_dir, pdf_file)
    size_mb = os.path.getsize(pdf_path) / (1024 * 1024)
    doc = fitz.open(pdf_path)
    total_pages = len(doc)
    
    text_pages_count = 0
    scanned_pages_count = 0
    sample_page_titles = []
    
    for page_idx in range(total_pages):
        page = doc[page_idx]
        text = page.get_text().strip()
        lines = [line.strip() for line in text.split("\n") if line.strip()]
        if len(text) > 40:
            text_pages_count += 1
            if len(sample_page_titles) < 15 and lines:
                sample_page_titles.append((page_idx + 1, lines[0][:60]))
        else:
            # Check if has images
            images = page.get_images()
            scanned_pages_count += 1
            if len(sample_page_titles) < 15:
                img_desc = f"[Raster/Image: {len(images)} imgs]"
                sample_page_titles.append((page_idx + 1, img_desc))

    print(f"\n[PDF] {pdf_file}")
    print(f"  Size: {size_mb:.2f} MB | Total Pages: {total_pages}")
    print(f"  Text pages: {text_pages_count} | Scanned/Raster pages: {scanned_pages_count}")
    print(f"  Metadata: {doc.metadata}")
    print(f"  Sample page content preview:")
    for p_num, title in sample_page_titles[:10]:
        print(f"    Page {p_num:2d}: {title}")

# 2. Inspect DOCX
docx_files = [f for f in os.listdir(data_dir) if f.endswith(".docx")]
for docx_file in docx_files:
    docx_path = os.path.join(data_dir, docx_file)
    size_kb = os.path.getsize(docx_path) / 1024
    doc = Document(docx_path)
    paras = doc.paragraphs
    tables = doc.tables
    print(f"\n[DOCX] {docx_file}")
    print(f"  Size: {size_kb:.2f} KB | Paragraphs: {len(paras)} | Tables: {len(tables)}")
    headings = [p.text.strip() for p in paras if p.style.name.startswith("Heading") or (len(p.text.strip()) > 3 and p.runs and p.runs[0].bold)]
    print(f"  Key headings/bold lines ({len(headings)} found):")
    for h in headings[:15]:
        print(f"    - {h[:80]}")

# 3. Inspect JSON & GeoJSON
json_files = [f for f in os.listdir(data_dir) if f.endswith(".json") or f.endswith(".geojson")]
for jf in sorted(json_files):
    jpath = os.path.join(data_dir, jf)
    size_kb = os.path.getsize(jpath) / 1024
    with open(jpath, "r", encoding="utf-8") as f:
        data = json.load(f)
    if "type" in data and data["type"] == "FeatureCollection":
        feats = data.get("features", [])
        geom_types = set(f.get("geometry", {}).get("type", "None") for f in feats)
        props_keys = set()
        for f in feats[:10]:
            props_keys.update(f.get("properties", {}).keys())
        print(f"\n[GeoJSON] {jf}")
        print(f"  Size: {size_kb:.1f} KB | Features: {len(feats)} | Geometries: {geom_types}")
        print(f"  Properties sample keys: {sorted(list(props_keys))[:10]}")
    else:
        keys = list(data.keys()) if isinstance(data, dict) else f"Array of {len(data)}"
        print(f"\n[JSON] {jf}")
        print(f"  Size: {size_kb:.1f} KB | Keys / Schema: {keys}")
        if isinstance(data, dict) and "questions" in data:
            print(f"  Questions count: {len(data['questions'])}")
            formats = set(q.get("chart_type") or q.get("type") for q in data["questions"])
            print(f"  Question chart types: {formats}")
        elif isinstance(data, dict) and "layers" in data:
            print(f"  Layers in master: {list(data.get('layers', {}).keys())}")
            if "metadata" in data:
                print(f"  Master metadata: {data['metadata']}")

